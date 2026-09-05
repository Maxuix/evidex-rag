"""Frozen simple-format content parity through the real local converter."""

from io import BytesIO
import hashlib
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from rag_kb.adapters.parser.docling.parser import _DoclingRuntime
from rag_kb.domain import ParserLimits, ParserProfile, ParserSource

def corpus():
    yield ParserSource('plain.txt', 'text/plain', '第一章\n数字 3.14 与版本 v1.2.3\n\n证据保持完整。'.encode())
    yield ParserSource('guide.md', 'text/markdown', b'# Heading\n\nA **bold** paragraph.\n\n| Name | Value |\n| --- | --- |\n| Alpha | 42 |\n')
    yield ParserSource('guide.html', 'text/html', b'<html><body><h1>Heading</h1><p>Text 3.14</p><table><tr><td>Alpha</td><td>42</td></tr></table></body></html>')
    yield ParserSource('table.csv', 'text/csv', b'Name,Value\nAlpha,42\nBeta,3.14\n')
    doc = Document(); doc.add_heading('Heading', 1); doc.add_paragraph('Body 3.14')
    table = doc.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, (('Name','Value'),('Alpha','42'))):
        for cell, value in zip(row.cells, values): cell.text = value
    output = BytesIO(); doc.save(output)
    yield ParserSource('word.docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', output.getvalue())
    book = Workbook(); book.active.title = 'Results'; book.active.append(['Name','Value']); book.active.append(['Alpha',42]); book.active.append(['Beta',3.14])
    output = BytesIO(); book.save(output)
    yield ParserSource('book.xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', output.getvalue())
    deck=Presentation(); slide=deck.slides.add_slide(deck.slide_layouts[5]); slide.shapes.title.text='Heading'
    box=slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1)); box.text_frame.text='Body 3.14 and Alpha 42'
    output=BytesIO(); deck.save(output)
    yield ParserSource('slides.pptx', 'application/vnd.openxmlformats-officedocument.presentationml.presentation', output.getvalue())


# Baseline exports recorded before review 18 fixes, with the pinned Docling/core.
EXPECTED = {'plain.txt': 'f49a66178ba761b65e22f1b0a388e24ac571df803b5fc6f3c993d0d839f856b6', 'guide.md': '96318ba9ee2f8afc4f706b1f470648ac68e73261c572a3ba1b02e904013cb6bd', 'guide.html': '1d3d844c025336b3790084ed770787fef857dd4b3486f9b67497209fe42c28c5', 'table.csv': 'c1ba8fc8079154d1bd2835d2fad043c8d1f55982edd5d02a5a919f2e0c35dca1', 'word.docx': '8da5e1bd587341151703633e1c15703f2b08fa625434ad47ef51e6189a898eb6', 'book.xlsx': 'c1ba8fc8079154d1bd2835d2fad043c8d1f55982edd5d02a5a919f2e0c35dca1', 'slides.pptx': '380edcd3fa21e62b10f329c630f618818980079038ad338e3e28756336057565'}

@pytest.mark.parametrize("source", list(corpus()), ids=lambda source: source.original_filename)
def test_real_simple_conversion_keeps_baseline_content(source, tmp_path):
    runtime = _DoclingRuntime(
        ParserLimits(), artifacts_path=tmp_path,
        artifact_manifest_path=tmp_path / "not-used-by-simple-pipelines.json",
    )
    document = runtime.convert(source, ParserProfile.DOCLING_TEXT_LOCAL_V4)
    content_hash = hashlib.sha256(document.export_to_markdown().encode()).hexdigest()
    assert content_hash == EXPECTED[source.original_filename]
