"""Small adversarial and evidence-fidelity regressions; no model calls."""

from dataclasses import replace
from io import BytesIO
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from PIL import Image
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import (
    ArrayObject, DecodedStreamObject, DictionaryObject, NameObject, NumberObject,
)
from docling_core.types.doc import (
    BoundingBox, DocItemLabel, DoclingDocument, ImageRef, ProvenanceItem,
    Size, TableCell, TableData,
)
from docling_core.types.doc.common.origin import DocumentOrigin

from rag_kb.adapters.parser.docling.parser import (
    _preflight_conversion_source, _validate_document,
)
from rag_kb.adapters.parser.scanned_pages import scanned_surfaces
from rag_kb.document_processing.docling.assets import extract_docling_assets
from rag_kb.domain import (
    AdmissionLimits, FileAdmissionError, ParserExecutionError, ParserLimits, ParserSource,
)
from rag_kb.services.admission import FileAdmissionService


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def pdf_source(*, stamp=False, nested=False, image_edge=100):
    writer = PdfWriter()
    page = writer.add_blank_page(width=100, height=100)
    bitmap = DecodedStreamObject()
    bitmap.set_data(b"\xff\xff\xff")
    bitmap.update({NameObject("/Type"): NameObject("/XObject"), NameObject("/Subtype"): NameObject("/Image"), NameObject("/Width"): NumberObject(1), NameObject("/Height"): NumberObject(1), NameObject("/ColorSpace"): NameObject("/DeviceRGB"), NameObject("/BitsPerComponent"): NumberObject(8)})
    objects = DictionaryObject({NameObject("/Im0"): writer._add_object(bitmap)})
    paint = f"q {image_edge} 0 0 {image_edge} 0 0 cm /Im0 Do Q".encode()
    if nested:
        form = DecodedStreamObject()
        form.set_data(paint)
        form.update({NameObject("/Type"): NameObject("/XObject"), NameObject("/Subtype"): NameObject("/Form"), NameObject("/BBox"): ArrayObject([NumberObject(n) for n in (0, 0, 100, 100)]), NameObject("/Resources"): DictionaryObject({NameObject("/XObject"): objects})})
        objects = DictionaryObject({NameObject("/Fm0"): writer._add_object(form)})
        paint = b"/Fm0 Do"
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/XObject"): objects, NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    content = DecodedStreamObject()
    content.set_data(paint + (b" BT /F1 8 Tf 1 1 Td (1) Tj ET" if stamp else b""))
    page[NameObject("/Contents")] = writer._add_object(content)
    output = BytesIO()
    writer.write(output)
    return ParserSource("scan.pdf", "application/pdf", output.getvalue())


def page_document(count=1, edge=200):
    doc = DoclingDocument(name="safe-fixture")
    doc.origin = DocumentOrigin(mimetype="application/pdf", binary_hash=1, filename="fixture.pdf")
    doc.add_text(label=DocItemLabel.TEXT, text="body")
    for number in range(1, count + 1):
        doc.add_page(page_no=number, size=Size(width=edge, height=edge), image=ImageRef.from_pil(Image.new("RGB", (edge, edge)), dpi=72))
    return doc


def xlsx_source(*, far_cell="B2", merge=None):
    book = Workbook()
    book.active["A1"] = "start"
    book.active[far_cell] = "end"
    content = BytesIO()
    book.save(content)
    if merge:
        # Inject XML, avoiding openpyxl materializing the giant merge itself.
        output = BytesIO()
        with ZipFile(BytesIO(content.getvalue())) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
            for item in source.infolist():
                payload = source.read(item)
                if item.filename == "xl/worksheets/sheet1.xml":
                    payload = payload.replace(b"</worksheet>", f'<mergeCells><mergeCell ref="{merge}"/></mergeCells></worksheet>'.encode())
                target.writestr(item.filename, payload)
        content = output
    return ParserSource("sparse.xlsx", XLSX, content.getvalue())


@pytest.mark.parametrize("source", [xlsx_source(far_cell="XFD1048576"), xlsx_source(merge="A1:XFD1048576")])
def test_xlsx_expansion_is_rejected_before_converter(source):
    with pytest.raises(FileAdmissionError):
        FileAdmissionService(AdmissionLimits()).validate(BytesIO(source.content), original_filename=source.original_filename, media_type=source.media_type)
    with pytest.raises(ParserExecutionError):
        _preflight_conversion_source(source, ParserLimits())


def test_small_xlsx_is_still_admitted():
    source = xlsx_source()
    FileAdmissionService(AdmissionLimits()).validate(BytesIO(source.content), original_filename=source.original_filename, media_type=source.media_type)
    _preflight_conversion_source(source, ParserLimits())


@pytest.mark.parametrize("stamp,nested", [(False, False), (True, False), (False, True), (True, True)])
def test_scanned_evidence_survives_stamp_and_form(stamp, nested):
    assert scanned_surfaces(pdf_source(stamp=stamp, nested=nested)) == frozenset({1})


def test_native_text_with_small_logo_needs_no_page_asset():
    assert scanned_surfaces(pdf_source(stamp=True, image_edge=10)) == frozenset()


def test_parser_checks_total_pixels():
    with pytest.raises(ParserExecutionError) as caught:
        _validate_document(page_document(count=2, edge=4), replace(ParserLimits(), max_total_image_pixels=20))
    assert caught.value.diagnostic["limit_name"] == "max_total_image_pixels"


def test_assets_check_total_pixels_before_decoding():
    document = page_document(count=2, edge=4)
    with patch.object(Image, "open", side_effect=AssertionError("pixel budget must reject before decoding")):
        with pytest.raises(ParserExecutionError) as caught:
            extract_docling_assets(document, replace(ParserLimits(), max_total_image_pixels=20), page_image_surfaces=frozenset({1, 2}))
    assert caught.value.diagnostic["limit_name"] == "max_total_image_pixels"


def test_table_is_cropped_from_page_image():
    document = page_document()
    table = document.add_table(data=TableData(num_rows=1, num_cols=1, table_cells=[TableCell(text="42", start_row_offset_idx=0, end_row_offset_idx=1, start_col_offset_idx=0, end_col_offset_idx=1)]), prov=ProvenanceItem(page_no=1, charspan=(0, 0), bbox=BoundingBox(l=0, t=0, r=100, b=100)))
    assert table.image is None
    assets = extract_docling_assets(document, page_image_surfaces=frozenset())
    assert [(a.kind, a.width, a.height) for a in assets] == [("table_image", 100, 100)]


def test_pdf_probe_does_not_extract_text():
    from pypdf._page import PageObject
    with patch.object(PageObject, "extract_text", side_effect=AssertionError("text extraction must not run")):
        assert scanned_surfaces(pdf_source(stamp=True, nested=True)) == frozenset({1})


def test_pdf_operator_and_decompression_budgets():
    from pypdf import PdfReader
    with pytest.raises(ParserExecutionError) as caught:
        scanned_surfaces(pdf_source(), replace(ParserLimits(), max_pdf_operators=1))
    assert caught.value.diagnostic["limit_name"] == "max_pdf_operators"
    reader = PdfReader(BytesIO(pdf_source().content))
    writer = PdfWriter()
    writer.add_page(reader.pages[0])
    content = DecodedStreamObject()
    content.set_data(b"q Q " * 1000)
    writer.pages[0][NameObject("/Contents")] = writer._add_object(content.flate_encode())
    output = BytesIO()
    writer.write(output)
    with pytest.raises(ParserExecutionError) as caught:
        scanned_surfaces(ParserSource("limited.pdf", "application/pdf", output.getvalue()), replace(ParserLimits(), max_pdf_content_bytes=100))
    assert caught.value.diagnostic["limit_name"] == "max_pdf_content_bytes"


def test_pdf_cyclic_form_is_rejected():
    from pypdf import PdfReader
    reader = PdfReader(BytesIO(pdf_source(nested=True).content))
    writer = PdfWriter()
    writer.add_page(reader.pages[0])
    form_ref = writer.pages[0]["/Resources"]["/XObject"]["/Fm0"]
    form = form_ref.get_object()
    form[NameObject("/Resources")] = DictionaryObject({NameObject("/XObject"): DictionaryObject({NameObject("/Loop"): form_ref.indirect_reference})})
    form.set_data(b"/Loop Do")
    output = BytesIO()
    writer.write(output)
    with pytest.raises(ParserExecutionError):
        scanned_surfaces(ParserSource("cycle.pdf", "application/pdf", output.getvalue()))


def test_image_header_must_match_declared_size():
    document = page_document(edge=100)
    document.pages[1].image.size = Size(width=4, height=4)
    with patch.object(Image.Image, "load", side_effect=AssertionError("no pixel decoding")) as load:
        with pytest.raises(ParserExecutionError) as caught:
            extract_docling_assets(document, page_image_surfaces=frozenset({1}))
        assert not load.called
    assert caught.value.diagnostic["check"] == "docling_image_header"


def test_table_crop_cannot_allocate_outside_pixel_budget():
    document = page_document()
    table = document.add_table(data=TableData(num_rows=0, num_cols=0, table_cells=[]), prov=ProvenanceItem(page_no=1, charspan=(0, 0), bbox=BoundingBox(l=0, t=0, r=10000, b=10000)))
    with patch.object(type(table), "get_image", side_effect=AssertionError("crop must be bounded first")):
        with pytest.raises(ParserExecutionError) as caught:
            extract_docling_assets(document, page_image_surfaces=frozenset())
    assert caught.value.diagnostic["limit_name"] == "max_image_pixels"


def _rewrite_package(source, change):
    output = BytesIO()
    with ZipFile(BytesIO(source.content)) as archive, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for item in archive.infolist():
            name, content = change(item.filename, archive.read(item))
            target.writestr(name, content)
    return ParserSource(source.original_filename, source.media_type, output.getvalue())


def test_xlsx_false_dimension_and_nonstandard_part_cannot_hide_span():
    import re
    def change(name, content):
        if name == "xl/_rels/workbook.xml.rels":
            content = content.replace(b"/xl/worksheets/sheet1.xml", b"/xl/alternate.dat")
        if name == "xl/worksheets/sheet1.xml":
            name = "xl/alternate.dat"
            content = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1"', content)
        return name, content
    source = _rewrite_package(xlsx_source(far_cell="XFD1048576"), change)
    with pytest.raises(ParserExecutionError):
        _preflight_conversion_source(source, ParserLimits())


def test_xlsx_dtd_is_rejected():
    def change(name, content):
        if name == "xl/worksheets/sheet1.xml":
            content = b'<!DOCTYPE worksheet [<!ENTITY x "expanded">]>' + content
        return name, content
    with pytest.raises(ParserExecutionError):
        _preflight_conversion_source(_rewrite_package(xlsx_source(), change), ParserLimits())


def test_xlsx_combines_spans_across_sheets():
    book = Workbook()
    book.active["A1"] = "one"
    book.active["D4"] = "two"
    sheet = book.create_sheet("two")
    sheet["A1"] = "one"
    sheet["D4"] = "two"
    content = BytesIO()
    book.save(content)
    with pytest.raises(ParserExecutionError) as caught:
        _preflight_conversion_source(ParserSource("two.xlsx", XLSX, content.getvalue()), replace(ParserLimits(), max_xlsx_cells=20))
    assert caught.value.diagnostic["limit_name"] == "max_xlsx_cells"


def test_checkpoint_rejects_combined_usage_before_writing(tmp_path):
    from uuid import uuid4
    from rag_kb.adapters.parser.docling.parser import DoclingParser
    from rag_kb.domain import ParserProfile
    parser = DoclingParser(replace(ParserLimits(), pdf_segment_pages=1, max_extracted_characters=6), artifacts_path=tmp_path, artifact_manifest_path=tmp_path / "unused.json", checkpoint_root=tmp_path)
    key = str(uuid4())
    state = {"stage_pages": {}, "ocr_pages": 0, "ocr_regions": 0, "table_candidates": 0, "child_peak_rss_bytes": None}
    try:
        checkpoint = parser._load_or_create_checkpoint(key, pdf_source(), ParserProfile.DOCLING_TEXT_LOCAL_V4, 2)
        first = page_document()
        checkpoint = parser._complete_segment(key, checkpoint, 0, first, 1, state)
        second = page_document()
        page = second.pages.pop(1)
        page.page_no = 2
        second.pages[2] = page
        with pytest.raises(ParserExecutionError) as caught:
            parser._complete_segment(key, checkpoint, 1, second, 1, state)
        assert caught.value.diagnostic["limit_name"] == "max_extracted_characters"
        assert not (parser._checkpoint_directory(key) / "pages-2-2.json").exists()
    finally:
        parser.close()


def test_legacy_checkpoint_overflow_is_detected_before_concatenation(tmp_path):
    from uuid import uuid4
    from rag_kb.adapters.parser.docling.parser import DoclingParser
    from rag_kb.domain import ParserProfile
    parser = DoclingParser(replace(ParserLimits(), pdf_segment_pages=1), artifacts_path=tmp_path, artifact_manifest_path=tmp_path / "unused.json", checkpoint_root=tmp_path)
    key = str(uuid4())
    state = {"stage_pages": {}, "ocr_pages": 0, "ocr_regions": 0, "table_candidates": 0, "child_peak_rss_bytes": None}
    try:
        checkpoint = parser._load_or_create_checkpoint(key, pdf_source(), ParserProfile.DOCLING_TEXT_LOCAL_V4, 2)
        for index in range(2):
            document = page_document()
            page = document.pages.pop(1)
            page.page_no = index + 1
            document.pages[index + 1] = page
            checkpoint = parser._complete_segment(key, checkpoint, index, document, 1, state)
        for segment in checkpoint["segments"]:
            segment.pop("usage")
        parser._limits = replace(parser._limits, max_extracted_characters=6)
        with patch.object(DoclingDocument, "concatenate", side_effect=AssertionError("must reject before concatenate")):
            with pytest.raises(ParserExecutionError) as caught:
                parser._assemble_checkpoint(key, checkpoint)
        assert caught.value.diagnostic["limit_name"] == "max_extracted_characters"
    finally:
        parser.close()


@pytest.mark.parametrize("platform,multiplier", [("darwin", 1), ("linux", 1024)])
def test_rss_units_are_correct(platform, multiplier):
    from types import SimpleNamespace
    from rag_kb.adapters.parser.docling import progress_pipeline
    with patch.object(progress_pipeline.sys, "platform", platform), patch.object(progress_pipeline.resource, "getrusage", return_value=SimpleNamespace(ru_maxrss=123)):
        assert progress_pipeline._child_peak_rss_bytes() == 123 * multiplier


def test_checkpoint_assembly_does_not_retain_all_segment_documents(tmp_path):
    import weakref
    from uuid import uuid4
    from rag_kb.adapters.parser.docling.parser import DoclingParser
    from rag_kb.domain import ParserProfile
    parser = DoclingParser(replace(ParserLimits(), pdf_segment_pages=1), artifacts_path=tmp_path, artifact_manifest_path=tmp_path / "unused.json", checkpoint_root=tmp_path)
    key = str(uuid4())
    state = {"stage_pages": {}, "ocr_pages": 0, "ocr_regions": 0, "table_candidates": 0, "child_peak_rss_bytes": None}
    references = []
    observed = []
    original = DoclingDocument.model_validate_json

    def track(payload, *args, **kwargs):
        document = original(payload, *args, **kwargs)
        references.append(weakref.ref(document))
        observed.append(sum(reference() is not None for reference in references))
        return document

    try:
        checkpoint = parser._load_or_create_checkpoint(key, pdf_source(), ParserProfile.DOCLING_TEXT_LOCAL_V4, 10)
        for index in range(10):
            document = page_document(edge=4)
            page = document.pages.pop(1)
            page.page_no = index + 1
            document.pages[index + 1] = page
            checkpoint = parser._complete_segment(key, checkpoint, index, document, 1, state)
        del document
        with patch.object(DoclingDocument, "model_validate_json", side_effect=track):
            merged = parser._assemble_checkpoint(key, checkpoint)
        assert len(merged.pages) == 10
        assert len(references) == 10
        assert max(observed) <= 2
    finally:
        parser.close()
