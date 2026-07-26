from __future__ import annotations

import unittest
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook

from rag_kb.adapters.parser.ooxml_metadata import worksheet_labels
from rag_kb.document_processing.docling.provenance import (
    ItemSurface,
    aggregate_provenance,
    surface_kind,
    surface_location,
)
from rag_kb.services.admission import FileAdmissionService
from rag_kb.domain import (
    AdmissionLimits,
    ErrorCode,
    FileAdmissionError,
    ParserExecutionError,
    ParserSource,
)


OOXML = "application/vnd.openxmlformats-officedocument"
XLSX_MEDIA_TYPE = f"{OOXML}.spreadsheetml.sheet"
PPTX_MEDIA_TYPE = f"{OOXML}.presentationml.presentation"
DOCX_MEDIA_TYPE = f"{OOXML}.wordprocessingml.document"


def workbook(sheets: tuple[tuple[str, tuple[tuple[str, ...], ...]], ...]) -> bytes:
    book = Workbook()
    book.remove(book.active)
    for title, rows in sheets:
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(list(row))
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def package(parts: dict[str, bytes]) -> bytes:
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        for name, payload in parts.items():
            archive.writestr(name, payload)
    return buffer.getvalue()


def source(name: str, content: bytes) -> ParserSource:
    return ParserSource(
        original_filename=name, media_type=XLSX_MEDIA_TYPE, content=content
    )


class AdmissionFormatTests(unittest.TestCase):
    def setUp(self) -> None:
        self.service = FileAdmissionService(AdmissionLimits())

    def admit(self, name: str, media_type: str, content: bytes):
        return self.service.validate(
            BytesIO(content), original_filename=name, media_type=media_type
        )

    def test_the_new_text_formats_are_admitted_as_utf8(self) -> None:
        cases = (
            ("page.html", "text/html", b"<html><body><p>hi</p></body></html>"),
            ("data.csv", "text/csv", b"a,b\n1,2\n"),
        )
        for name, media_type, content in cases:
            with self.subTest(name=name):
                admitted = self.admit(name, media_type, content)
                self.assertEqual(admitted.media_type, media_type)
                self.assertIsNotNone(admitted.line_count)

    def test_new_text_formats_reject_invalid_utf8(self) -> None:
        with self.assertRaises(FileAdmissionError) as raised:
            self.admit("page.html", "text/html", b"\xff\xfe not utf8")
        self.assertEqual(raised.exception.code, ErrorCode.FILE_INVALID_UTF8)

    def test_every_ooxml_family_requires_its_own_content_part(self) -> None:
        cases = (
            ("book.xlsx", XLSX_MEDIA_TYPE, "xl/workbook.xml"),
            ("deck.pptx", PPTX_MEDIA_TYPE, "ppt/presentation.xml"),
            ("text.docx", DOCX_MEDIA_TYPE, "word/document.xml"),
        )
        for name, media_type, required in cases:
            with self.subTest(name=name):
                self.admit(
                    name,
                    media_type,
                    package({"[Content_Types].xml": b"<Types/>", required: b"<x/>"}),
                )
                with self.assertRaises(FileAdmissionError) as raised:
                    self.admit(
                        name,
                        media_type,
                        package({"[Content_Types].xml": b"<Types/>"}),
                    )
                self.assertEqual(
                    raised.exception.code, ErrorCode.FILE_CONTENT_INVALID
                )

    def test_a_renamed_package_of_another_family_is_rejected(self) -> None:
        with self.assertRaises(FileAdmissionError) as raised:
            self.admit(
                "deck.pptx",
                PPTX_MEDIA_TYPE,
                package(
                    {"[Content_Types].xml": b"<Types/>", "word/document.xml": b"<x/>"}
                ),
            )
        self.assertEqual(raised.exception.code, ErrorCode.FILE_CONTENT_INVALID)

    def test_path_traversal_entries_are_rejected_for_every_family(self) -> None:
        for name, media_type, required in (
            ("book.xlsx", XLSX_MEDIA_TYPE, "xl/workbook.xml"),
            ("deck.pptx", PPTX_MEDIA_TYPE, "ppt/presentation.xml"),
        ):
            with self.subTest(name=name):
                with self.assertRaises(FileAdmissionError) as raised:
                    self.admit(
                        name,
                        media_type,
                        package(
                            {
                                "[Content_Types].xml": b"<Types/>",
                                required: b"<x/>",
                                "../escape.xml": b"<x/>",
                            }
                        ),
                    )
                self.assertEqual(
                    raised.exception.code, ErrorCode.FILE_CONTENT_INVALID
                )

    def test_the_archive_entry_budget_is_shared(self) -> None:
        service = FileAdmissionService(AdmissionLimits(max_archive_entries=2))
        parts = {
            "[Content_Types].xml": b"<Types/>",
            "xl/workbook.xml": b"<x/>",
            "xl/extra.xml": b"<x/>",
        }
        with self.assertRaises(FileAdmissionError) as raised:
            service.validate(
                BytesIO(package(parts)),
                original_filename="book.xlsx",
                media_type=XLSX_MEDIA_TYPE,
            )
        self.assertEqual(
            raised.exception.code, ErrorCode.FILE_ARCHIVE_LIMIT_EXCEEDED
        )

    def test_extension_and_media_type_must_agree(self) -> None:
        with self.assertRaises(FileAdmissionError) as raised:
            self.admit("data.csv", "text/html", b"a,b\n")
        self.assertEqual(raised.exception.code, ErrorCode.FILE_MEDIA_TYPE_MISMATCH)


class WorksheetLabelTests(unittest.TestCase):
    def test_sheet_names_are_recovered_in_workbook_order(self) -> None:
        content = workbook(
            (
                ("Latency", (("Region", "Value"), ("Amber", "22"))),
                ("Costs", (("Item", "Cost"), ("Storage", "120"))),
            )
        )

        self.assertEqual(
            worksheet_labels(source("book.xlsx", content)),
            {1: "Latency", 2: "Costs"},
        )

    def test_non_spreadsheets_supply_no_labels(self) -> None:
        self.assertEqual(worksheet_labels(source("notes.txt", b"plain")), {})

    def test_a_corrupt_workbook_part_fails_closed(self) -> None:
        content = package(
            {"[Content_Types].xml": b"<Types/>", "xl/workbook.xml": b"<not-xml"}
        )
        with self.assertRaises(ParserExecutionError) as raised:
            worksheet_labels(source("book.xlsx", content))
        self.assertEqual(raised.exception.code, ErrorCode.FILE_CONTENT_INVALID)

    def test_a_workbook_without_named_sheets_fails_closed(self) -> None:
        content = package(
            {"[Content_Types].xml": b"<Types/>", "xl/workbook.xml": b"<workbook/>"}
        )
        with self.assertRaises(ParserExecutionError) as raised:
            worksheet_labels(source("book.xlsx", content))
        self.assertEqual(raised.exception.code, ErrorCode.FILE_CONTENT_INVALID)

    def test_repeated_reads_are_identical(self) -> None:
        content = workbook((("Only", (("a",),)),))
        probe = source("book.xlsx", content)

        self.assertEqual(worksheet_labels(probe), worksheet_labels(probe))


class SurfaceLabelProjectionTests(unittest.TestCase):
    def test_a_single_sheet_chunk_is_named(self) -> None:
        projection = aggregate_provenance(
            ((ItemSurface(kind="sheet", ordinal=2),),),
            ("#/tables/1",),
            max_metadata_bytes=65_536,
            surface_labels={1: "Latency", 2: "Costs"},
        )

        self.assertEqual(projection["surface_type"], "sheet")
        self.assertEqual(projection["surface_start"], 2)
        self.assertEqual(projection["surface_label"], "Costs")

    def test_a_chunk_spanning_sheets_lists_their_names(self) -> None:
        projection = aggregate_provenance(
            (
                (ItemSurface(kind="sheet", ordinal=1),),
                (ItemSurface(kind="sheet", ordinal=2),),
            ),
            ("#/tables/0", "#/tables/1"),
            max_metadata_bytes=65_536,
            surface_labels={1: "Latency", 2: "Costs"},
        )

        self.assertEqual(projection["surface_labels"], ["Latency", "Costs"])
        self.assertNotIn("surface_label", projection)

    def test_unnamed_surfaces_stay_unlabelled(self) -> None:
        projection = aggregate_provenance(
            ((ItemSurface(kind="page", ordinal=3),),),
            ("#/texts/0",),
            max_metadata_bytes=65_536,
            surface_labels={1: "Latency"},
        )

        self.assertNotIn("surface_label", projection)
        self.assertNotIn("surface_labels", projection)

    def test_whole_surface_assets_carry_the_name(self) -> None:
        self.assertEqual(
            surface_location("sheet", 1, {1: "Latency"})["surface_label"], "Latency"
        )
        self.assertNotIn("surface_label", surface_location("page", 1))


class SurfaceKindTests(unittest.TestCase):
    def test_presentation_packages_map_to_the_slide_surface(self) -> None:
        from docling_core.types.doc import DoclingDocument
        from docling_core.types.doc.common.origin import DocumentOrigin

        for mimetype, expected in (
            (f"{OOXML}.presentationml.presentation", "slide"),
            # Docling reports the template mimetype for decks built from one.
            (f"{OOXML}.presentationml.template", "slide"),
            (f"{OOXML}.spreadsheetml.sheet", "sheet"),
            ("text/html", "logical"),
            ("text/csv", "logical"),
        ):
            with self.subTest(mimetype=mimetype):
                document = DoclingDocument(name="probe")
                document.origin = DocumentOrigin(
                    mimetype=mimetype, binary_hash=1, filename="probe"
                )
                self.assertEqual(surface_kind(document), expected)


if __name__ == "__main__":
    unittest.main()
