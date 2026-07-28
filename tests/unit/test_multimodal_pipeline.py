from __future__ import annotations

import hashlib
import math
import tempfile
import unittest
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from uuid import UUID, uuid4

import httpx
from docx import Document
from PIL import Image
from pypdf import PdfWriter

from rag_kb.adapters import LocalIndexAssetStore, TongyiVisionEmbeddingAdapter
from apps.api.routers.assets import read_index_asset
from rag_kb.auth import AuthContext, SingleWorkspaceAccessPolicy
from rag_kb.adapters.parser.scanned_pages import scanned_surfaces
from rag_kb.domain import (
    ContentModality,
    EmbeddingSpaceDefinition,
    ErrorCode,
    ImageEmbeddingInput,
    IndexAssetIdentity,
    IndexAssetSnapshot,
    IndexingExecutionError,
    ParserExecutionError,
    ParserLimits,
    ParserSource,
    VectorSearchHit,
)
from rag_kb.retrieval.fusion import reciprocal_rank_fusion
from rag_kb.services import IndexAssetService


WORKSPACE = UUID("01900000-0000-7000-8000-000000000901")
INDEXED_VERSION = UUID("01900000-0000-7000-8000-000000000902")


def _png(width: int = 120, height: int = 80) -> bytes:
    output = BytesIO()
    Image.new("RGB", (width, height), (32, 96, 160)).save(output, format="PNG")
    return output.getvalue()


def _docx() -> bytes:
    document = Document()
    document.add_heading("Architecture", level=1)
    document.add_paragraph("The body text remains an independent text unit.")
    document.add_picture(BytesIO(_png()))
    document.add_paragraph("Figure 1: service topology", style="Caption")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Service"
    table.cell(0, 1).text = "Owner"
    table.cell(1, 0).text = "API"
    table.cell(1, 1).text = "Platform"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


class _RawElement:
    def __init__(self, category: str, text: str, metadata: dict[str, object]) -> None:
        self.category = category
        self._text = text
        self.metadata = _RawMetadata(metadata)

    def __str__(self) -> str:
        return self._text


class _RawMetadata:
    def __init__(self, value: dict[str, object]) -> None:
        self._value = value

    def to_dict(self) -> dict[str, object]:
        return dict(self._value)


def _space() -> EmbeddingSpaceDefinition:
    return EmbeddingSpaceDefinition(
        provider_identity="alibaba-cloud-model-studio-qwen",
        requested_model="tongyi-embedding-vision-flash-2026-03-06",
        resolved_model="tongyi-embedding-vision-flash-2026-03-06",
        model_version="tongyi-embedding-vision-flash-2026-03-06",
        dimension=768,
        distance_metric="cosine",
        vector_data_type="float32",
        normalization="l2",
        endpoint_identity="alibaba-model-studio-beijing-multimodal-embedding",
        deployment_revision=None,
        configuration_fingerprint="sha256:test-configuration",
        tokenizer_fingerprint=None,
        compatibility_fingerprint="sha256:test-multimodal",
    )


class ScannedSurfaceTests(unittest.TestCase):
    def test_text_layer_probe_stops_after_first_non_whitespace_fragment(self) -> None:
        class TextPage:
            def __init__(self) -> None:
                self.visited: list[str] = []
                self.swallowed_control_flow = False
                self.reached_later_text = False

            def extract_text(self, *, visitor_text) -> str:
                visitor_text(" \n", None, None, None, None)
                self.visited.append("whitespace")
                try:
                    visitor_text("native text", None, None, None, None)
                except Exception:
                    self.swallowed_control_flow = True
                self.reached_later_text = True
                visitor_text("later text", None, None, None, None)
                return "native text later text"

        page = TextPage()
        with patch(
            "rag_kb.adapters.parser.scanned_pages.PdfReader",
            return_value=SimpleNamespace(pages=(page,)),
        ):
            result = scanned_surfaces(_pdf_source(b"synthetic"))

        self.assertEqual(result, frozenset())
        self.assertEqual(page.visited, ["whitespace"])
        self.assertFalse(page.swallowed_control_flow)
        self.assertFalse(page.reached_later_text)

    def test_raster_only_pdf_page_is_detected_without_provider_or_ocr(self) -> None:
        scanned = BytesIO()
        Image.new("RGB", (120, 80), "white").save(scanned, format="PDF")

        self.assertEqual(
            scanned_surfaces(_pdf_source(scanned.getvalue())), frozenset({1})
        )

        blank = BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        writer.write(blank)

        self.assertEqual(scanned_surfaces(_pdf_source(blank.getvalue())), frozenset())

    def test_non_pdf_sources_report_no_scanned_surface(self) -> None:
        self.assertEqual(
            scanned_surfaces(
                ParserSource(
                    original_filename="notes.txt",
                    media_type="text/plain",
                    content=b"plain",
                )
            ),
            frozenset(),
        )

    def test_corrupt_pdf_bytes_fail_closed(self) -> None:
        with self.assertRaises(ParserExecutionError) as raised:
            scanned_surfaces(_pdf_source(b"%PDF-1.7 truncated"))
        self.assertEqual(raised.exception.code, ErrorCode.PARSER_OUTPUT_INVALID)


def _pdf_source(content: bytes) -> ParserSource:
    return ParserSource(
        original_filename="probe.pdf",
        media_type="application/pdf",
        content=content,
    )


class IndexAssetStoreTests(unittest.IsolatedAsyncioTestCase):
    async def test_atomic_idempotent_asset_round_trip_and_checksum_guard(self) -> None:
        with tempfile.TemporaryDirectory() as root:
            staging = Path(root) / "staging"
            final = Path(root) / "final"
            staging.mkdir()
            final.mkdir()
            store = LocalIndexAssetStore(staging, final)
            content = _png()
            checksum = hashlib.sha256(content).hexdigest()
            identity = IndexAssetIdentity(WORKSPACE, INDEXED_VERSION, checksum)

            await store.put(identity, content, checksum)
            await store.put(identity, content, checksum)
            self.assertEqual(await store.read(identity), content)
            self.assertEqual(store.parse_uri(identity.storage_uri), identity)
            with self.assertRaises(Exception):
                await store.put(identity, b"different", checksum)
            await store.delete(identity)
            with self.assertRaises(Exception):
                await store.read(identity)

    async def test_authorized_service_and_http_response_keep_binary_out_of_json(self) -> None:
        with tempfile.TemporaryDirectory() as root:
            staging = Path(root) / "staging"
            final = Path(root) / "final"
            staging.mkdir()
            final.mkdir()
            store = LocalIndexAssetStore(staging, final)
            content = _png()
            checksum = hashlib.sha256(content).hexdigest()
            identity = IndexAssetIdentity(WORKSPACE, INDEXED_VERSION, checksum)
            await store.put(identity, content, checksum)
            snapshot = IndexAssetSnapshot(
                id=uuid4(),
                workspace_id=WORKSPACE,
                kb_id=uuid4(),
                document_id=uuid4(),
                document_version_id=uuid4(),
                indexed_document_version_id=INDEXED_VERSION,
                storage_uri=identity.storage_uri,
                media_type="image/png",
                checksum_sha256=checksum,
            )
            service = IndexAssetService(
                _AssetUowFactory(snapshot),
                SingleWorkspaceAccessPolicy(WORKSPACE),
                store,
            )
            context = AuthContext("principal", "client", WORKSPACE)

            loaded = await service.read(context, snapshot.id)
            response = await read_index_asset(
                SimpleNamespace(
                    app=SimpleNamespace(
                        state=SimpleNamespace(
                            dependencies=SimpleNamespace(
                                index_asset_service=service
                            )
                        )
                    )
                ),
                snapshot.id,
                context,
            )

            self.assertEqual(loaded.content, content)
            self.assertEqual(response.body, content)
            self.assertEqual(response.headers["content-type"], "image/png")
            self.assertEqual(response.headers["etag"], f'"sha256:{checksum}"')
            self.assertEqual(response.headers["x-content-type-options"], "nosniff")
            self.assertEqual(response.headers["cache-control"], "private, max-age=300")


class MultimodalEmbeddingAdapterTests(unittest.IsolatedAsyncioTestCase):
    async def test_text_and_image_inputs_are_separate_and_response_is_validated(self) -> None:
        requests: list[dict] = []
        vector = [0.0] * 768
        vector[0] = 1.0

        async def handler(request: httpx.Request) -> httpx.Response:
            payload = __import__("json").loads(request.content)
            requests.append(payload)
            count = len(payload["input"]["contents"])
            return httpx.Response(
                200,
                json={
                    "output": {
                        "embeddings": [
                            {"index": index, "embedding": vector}
                            for index in range(count)
                        ]
                    }
                },
            )

        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            adapter = TongyiVisionEmbeddingAdapter(
                endpoint="https://provider.invalid/embeddings",
                api_key="test-only",
                embedding_space=_space(),
                max_batch_size=20,
                timeout_seconds=2,
                max_retries=0,
                max_concurrency=1,
                client=client,
            )
            text = await adapter.embed_texts(("diagram",))
            image = await adapter.embed_images(
                (ImageEmbeddingInput(_png(), "image/png", hashlib.sha256(_png()).hexdigest()),)
            )

        self.assertEqual(text.vectors[0][0], 1.0)
        self.assertEqual(image.vectors[0][0], 1.0)
        self.assertEqual(requests[0]["input"]["contents"], [{"text": "query: diagram"}])
        self.assertEqual(
            requests[0]["parameters"],
            {"dimension": 768, "output_type": "dense", "res_level": 1},
        )
        self.assertTrue(requests[1]["input"]["contents"][0]["image"].startswith("data:image/png;base64,"))

    async def test_non_normalized_provider_vector_is_rejected(self) -> None:
        async def handler(_: httpx.Request) -> httpx.Response:
            return httpx.Response(
                200,
                json={"output": {"embeddings": [{"index": 0, "embedding": [1.0] * 768}]}},
            )

        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            adapter = TongyiVisionEmbeddingAdapter(
                endpoint="https://provider.invalid/embeddings",
                api_key="test-only",
                embedding_space=_space(),
                max_batch_size=20,
                timeout_seconds=2,
                max_retries=0,
                max_concurrency=1,
                client=client,
            )
            with self.assertRaises(IndexingExecutionError) as invalid:
                await adapter.embed_texts(("query",))
        self.assertEqual(invalid.exception.code, ErrorCode.EMBEDDING_RESPONSE_INVALID)


class RankFusionTests(unittest.TestCase):
    def test_rrf_deduplicates_evidence_group_and_preserves_lane_diagnostics(self) -> None:
        group = "asset-group"
        text = _hit(uuid4(), "caption_text", "image", group, 0.08, "caption")
        duplicate_text = _hit(uuid4(), "ocr_text", "image", group, 0.09, "ocr")
        native = _hit(uuid4(), "native_image", "image", group, 0.04, "")

        fused = reciprocal_rank_fusion(
            (text, duplicate_text), (native,), rrf_k=60, top_k=10
        )

        self.assertEqual(len(fused), 1)
        self.assertEqual(fused[0].text_rank, 1)
        self.assertEqual(fused[0].cross_modal_rank, 1)
        self.assertEqual(
            set(fused[0].matched_representations),
            {"caption_text", "ocr_text", "native_image"},
        )
        self.assertTrue(math.isclose(fused[0].score, 2 / 61, abs_tol=1e-12))

    def test_rrf_keeps_same_raw_group_separate_across_index_targets(self) -> None:
        group = "docling-table-group"
        text = _hit(uuid4(), "table_text", "table", group, 0.08, "table A")
        other_target = uuid4()
        native = replace(
            _hit(uuid4(), "table_image", "table", group, 0.04, "table B"),
            indexed_document_version_id=other_target,
            document_id=uuid4(),
            document_version_id=uuid4(),
        )

        fused = reciprocal_rank_fusion((text,), (native,), rrf_k=60, top_k=10)

        self.assertEqual(len(fused), 2)
        self.assertEqual(
            {item.hit.indexed_document_version_id for item in fused},
            {INDEXED_VERSION, other_target},
        )
        self.assertEqual(
            {item.matched_representations for item in fused},
            {("table_text",), ("table_image",)},
        )
        self.assertTrue(
            all(math.isclose(item.score, 1 / 61, abs_tol=1e-12) for item in fused)
        )


def _hit(
    chunk_id: UUID,
    representation: str,
    modality: str,
    group: str,
    distance: float,
    text: str,
) -> VectorSearchHit:
    return VectorSearchHit(
        workspace_id=WORKSPACE,
        knowledge_base_id=uuid4(),
        index_revision_id=uuid4(),
        index_chunk_id=chunk_id,
        indexed_document_version_id=INDEXED_VERSION,
        document_id=uuid4(),
        document_version_id=uuid4(),
        ordinal=0,
        text=text,
        source_location={},
        hierarchy={},
        source_metadata={},
        cosine_distance=distance,
        build_status="ready",
        serving_status="serving",
        is_current_serving_version=True,
        modality=modality,
        evidence_group_key=group,
        representation_kind=representation,
    )


class _AssetRepository:
    def __init__(self, snapshot: IndexAssetSnapshot) -> None:
        self.snapshot = snapshot

    async def get_asset(self, asset_id: UUID):
        return self.snapshot if asset_id == self.snapshot.id else None


class _AssetUow:
    def __init__(self, snapshot: IndexAssetSnapshot) -> None:
        self.workspace_id = WORKSPACE
        self.indexing = _AssetRepository(snapshot)

    async def __aenter__(self):
        return self

    async def __aexit__(self, *args):
        return None

    async def commit(self):
        return None


class _AssetUowFactory:
    def __init__(self, snapshot: IndexAssetSnapshot) -> None:
        self.snapshot = snapshot

    def __call__(self, **kwargs):
        del kwargs
        return _AssetUow(self.snapshot)
