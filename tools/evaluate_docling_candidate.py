#!/usr/bin/env python3
"""Evaluate an isolated Docling candidate against generated multi-format fixtures."""

from __future__ import annotations

import argparse
from collections import Counter
import hashlib
from importlib.metadata import PackageNotFoundError, version
from io import BytesIO
import json
import os
from pathlib import Path
import platform
import resource
import subprocess
import sys
import tempfile
import time
from typing import Any


MARKERS = {
    "plain_text": ("TXT-ALPHA-101",),
    "markdown": ("MD-FIRST-201", "MD-SECOND-202"),
    "html": ("HTML-FIRST-301", "HTML-TABLE-303"),
    "csv": ("CSV-TABLE-404",),
    "text_pdf": ("PDF-TEXT-505",),
    "scanned_pdf": ("SCAN-OCR-606",),
    "docx": ("DOCX-BODY-707", "DOCX-CAPTION-708", "DOCX-TABLE-709"),
    "pptx": ("PPTX-TITLE-801", "PPTX-BODY-802", "PPTX-TABLE-803"),
    "xlsx": ("XLSX-SHEET-901", "XLSX-TABLE-902"),
}

REQUIRED_CONTENT_MARKERS = {
    **MARKERS,
    # Docling 2.114.0 preserves the XLSX table body but not the worksheet name.
    # Keep the sheet marker as a fidelity assertion without treating it as a
    # conversion failure.
    "xlsx": ("XLSX-TABLE-902",),
}

LANGCHAIN_EXPORT_MODES = (
    "markdown",
    "doc_chunks_hierarchical",
    "doc_chunks_hybrid",
)

CORPUS_FILENAMES = {
    "plain_text": "plain.txt",
    "markdown": "structured.md",
    "html": "structured.html",
    "csv": "structured.csv",
    "text_pdf": "text.pdf",
    "scanned_pdf": "scan.pdf",
    "docx": "structured.docx",
    "pptx": "structured.pptx",
    "xlsx": "structured.xlsx",
}

CURRENT_MEDIA_TYPES = {
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--project-python",
        type=Path,
        help="project interpreter used for the current-parser comparison",
    )
    parser.add_argument(
        "--cache-dir",
        type=Path,
        help="candidate model/cache root; defaults to a temporary directory",
    )
    parser.add_argument(
        "--output",
        type=Path,
        help="optional JSON report destination",
    )
    parser.add_argument(
        "--current-worker",
        type=Path,
        help=argparse.SUPPRESS,
    )
    arguments = parser.parse_args()
    if arguments.current_worker is not None:
        print(
            json.dumps(
                evaluate_current_parser(arguments.current_worker),
                ensure_ascii=False,
                sort_keys=True,
            )
        )
        return 0
    if arguments.project_python is None:
        parser.error("--project-python is required")
    project_python = arguments.project_python.absolute()
    if not project_python.is_file():
        parser.error(f"--project-python does not exist: {project_python}")
    with tempfile.TemporaryDirectory(prefix="rag-kb-docling-corpus-") as directory:
        corpus_root = Path(directory)
        corpus = generate_corpus(corpus_root)
        if arguments.cache_dir is None:
            with tempfile.TemporaryDirectory(
                prefix="rag-kb-docling-cache-"
            ) as cache_directory:
                report = evaluate(
                    corpus,
                    project_python=project_python,
                    cache_root=Path(cache_directory),
                )
        else:
            cache_root = arguments.cache_dir.resolve()
            cache_root.mkdir(parents=True, exist_ok=True)
            report = evaluate(
                corpus,
                project_python=project_python,
                cache_root=cache_root,
            )
    serialized = json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True)
    if arguments.output is not None:
        arguments.output.parent.mkdir(parents=True, exist_ok=True)
        arguments.output.write_text(serialized + "\n", encoding="utf-8")
        print(str(arguments.output.resolve()))
    else:
        print(serialized)
    return 0 if report["decision"]["status"] != "no_go" else 1


def evaluate(
    corpus: dict[str, Path],
    *,
    project_python: Path,
    cache_root: Path,
) -> dict[str, Any]:
    os.environ["HF_HOME"] = str(cache_root / "huggingface")
    os.environ["DOCLING_CACHE_DIR"] = str(cache_root / "docling")
    started = time.perf_counter()
    native = evaluate_docling_native(corpus)
    langchain = evaluate_langchain_adapter(corpus)
    current = run_current_worker(project_python, corpus)
    installation_bytes = directory_size(Path(sys.prefix))
    cache_bytes = directory_size(cache_root)
    candidate_success = all(
        item["success"]
        and item["required_content_complete"]
        and item["deterministic_projection"]
        for item in native["formats"].values()
    )
    adapter_success = all(
        all(
            item[mode]["success"]
            and item[mode]["required_content_complete"]
            and item[mode]["deterministic"]
            for mode in LANGCHAIN_EXPORT_MODES
        )
        for item in langchain["formats"].values()
    )
    structure_ready = all(
        (
            native["formats"][label]["table_count"] >= minimum_tables
            and native["formats"][label]["picture_count"] >= minimum_pictures
        )
        for label, minimum_tables, minimum_pictures in (
            ("html", 1, 0),
            ("csv", 1, 0),
            ("docx", 1, 1),
            ("pptx", 1, 1),
            ("xlsx", 1, 1),
        )
    )
    status = (
        "go"
        if candidate_success and adapter_success and structure_ready
        else "conditional_go"
        if candidate_success
        else "no_go"
    )
    return {
        "schema_version": 1,
        "candidate": {
            "docling": package_version("docling"),
            "docling_core": package_version("docling-core"),
            "langchain_docling": package_version("langchain-docling"),
            "langchain_core": package_version("langchain-core"),
            "python": platform.python_version(),
            "platform": platform.platform(),
            "isolated_environment_bytes": installation_bytes,
            "cache_bytes_after_evaluation": cache_bytes,
        },
        "corpus": {
            label: {
                "suffix": path.suffix.lower(),
                "bytes": path.stat().st_size,
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                "marker_count": len(MARKERS[label]),
            }
            for label, path in corpus.items()
        },
        "native_docling": native,
        "langchain_adapter": langchain,
        "current_application": current,
        "resource": {
            "elapsed_seconds": round(time.perf_counter() - started, 3),
            "peak_rss_bytes": peak_rss_bytes(),
        },
        "decision": {
            "status": status,
            "native_all_formats_passed": candidate_success,
            "langchain_exports_passed": adapter_success,
            "rich_structure_minimums_passed": structure_ready,
            "all_fidelity_markers_passed": all(
                item["markers_complete"] for item in native["formats"].values()
            ),
            "note": (
                "Generated fixtures establish candidate feasibility only; a go result "
                "still requires real enterprise-document shadow evaluation."
            ),
        },
    }


def generate_corpus(root: Path) -> dict[str, Path]:
    from docx import Document
    from docx.shared import Inches
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as WorksheetImage
    from PIL import Image, ImageDraw
    from pptx import Presentation
    from pptx.util import Inches as PresentationInches
    if __package__:
        from tools.evaluate_multimodal_real import (
            _architecture_image,
            _font,
            _write_pdf,
        )
    else:
        from evaluate_multimodal_real import _architecture_image, _font, _write_pdf

    image_path = root / "shared-figure.png"
    _architecture_image().save(image_path, "PNG")

    plain_text = root / "plain.txt"
    plain_text.write_text(
        "TXT-ALPHA-101\nA bounded plain-text parser fixture.", encoding="utf-8"
    )
    markdown = root / "structured.md"
    markdown.write_text(
        "# MD-FIRST-201\n\nBody evidence.\n\n## MD-SECOND-202\n\nFinal evidence.",
        encoding="utf-8",
    )
    html = root / "structured.html"
    html.write_text(
        "<html><body><h1>HTML-FIRST-301</h1>"
        "<table><tr><th>Key</th><th>Value</th></tr>"
        "<tr><td>HTML-TABLE-303</td><td>approved</td></tr></table>"
        "</body></html>",
        encoding="utf-8",
    )
    csv = root / "structured.csv"
    csv.write_text(
        "key,value\nCSV-TABLE-404,approved\n", encoding="utf-8"
    )

    text_pdf = root / "text.pdf"
    _write_pdf(
        text_pdf,
        (
            "PDF-TEXT-505",
            "The PDF contains native text plus a bounded figure.",
            "Figure 5. Shared parser architecture.",
        ),
        _architecture_image(),
        image_draws=((72, 280, 360, 220),),
    )
    scanned_pdf = root / "scan.pdf"
    scan = Image.new("RGB", (1500, 900), "white")
    drawing = ImageDraw.Draw(scan)
    drawing.text((160, 280), "SCAN-OCR-606", fill="black", font=_font(92))
    drawing.rectangle((100, 200, 1000, 500), outline="navy", width=8)
    scan.save(scanned_pdf, "PDF", resolution=150.0)

    docx_path = root / "structured.docx"
    document = Document()
    document.add_heading("DOCX-BODY-707", level=1)
    document.add_paragraph("Body before the embedded figure.")
    document.add_picture(str(image_path), width=Inches(3.0))
    document.add_paragraph("Figure 7. DOCX-CAPTION-708", style="Caption")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Key"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "DOCX-TABLE-709"
    table.cell(1, 1).text = "approved"
    document.save(docx_path)

    pptx_path = root / "structured.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "PPTX-TITLE-801"
    body = slide.shapes.add_textbox(
        PresentationInches(0.7),
        PresentationInches(1.2),
        PresentationInches(4.5),
        PresentationInches(0.7),
    )
    body.text_frame.text = "PPTX-BODY-802"
    slide.shapes.add_picture(
        str(image_path),
        PresentationInches(5.4),
        PresentationInches(1.0),
        width=PresentationInches(3.0),
    )
    table_shape = slide.shapes.add_table(
        2,
        2,
        PresentationInches(0.7),
        PresentationInches(2.4),
        PresentationInches(4.2),
        PresentationInches(1.5),
    )
    table_shape.table.cell(0, 0).text = "Key"
    table_shape.table.cell(0, 1).text = "Value"
    table_shape.table.cell(1, 0).text = "PPTX-TABLE-803"
    table_shape.table.cell(1, 1).text = "approved"
    presentation.save(pptx_path)

    xlsx_path = root / "structured.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "XLSX-SHEET-901"
    sheet.append(("Key", "Value"))
    sheet.append(("XLSX-TABLE-902", "approved"))
    worksheet_image = WorksheetImage(str(image_path))
    worksheet_image.width = 240
    worksheet_image.height = 160
    sheet.add_image(worksheet_image, "D2")
    workbook.save(xlsx_path)

    return {
        "plain_text": plain_text,
        "markdown": markdown,
        "html": html,
        "csv": csv,
        "text_pdf": text_pdf,
        "scanned_pdf": scanned_pdf,
        "docx": docx_path,
        "pptx": pptx_path,
        "xlsx": xlsx_path,
    }


def evaluate_docling_native(corpus: dict[str, Path]) -> dict[str, Any]:
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import PdfPipelineOptions
    from docling.document_converter import DocumentConverter, PdfFormatOption

    pdf_options = PdfPipelineOptions()
    pdf_options.do_ocr = True
    pdf_options.generate_page_images = True
    pdf_options.generate_picture_images = True
    converter = DocumentConverter(
        format_options={
            InputFormat.PDF: PdfFormatOption(pipeline_options=pdf_options)
        }
    )
    first: dict[str, Any] = {}
    second: dict[str, Any] = {}
    for destination in (first, second):
        for label, path in corpus.items():
            destination[label] = convert_native(converter, label, path)
    formats: dict[str, Any] = {}
    for label in corpus:
        initial = first[label]
        repeated = second[label]
        formats[label] = {
            **initial,
            "warm_elapsed_seconds": repeated["elapsed_seconds"],
            "deterministic_projection": (
                initial["projection_sha256"] == repeated["projection_sha256"]
            ),
            "deterministic_assets": (
                initial["asset_sha256"] == repeated["asset_sha256"]
            ),
        }
    return {
        "formats": formats,
        "cold_total_seconds": round(
            sum(item["elapsed_seconds"] for item in first.values()), 3
        ),
        "warm_total_seconds": round(
            sum(item["elapsed_seconds"] for item in second.values()), 3
        ),
    }


def convert_native(converter: Any, label: str, path: Path) -> dict[str, Any]:
    started = time.perf_counter()
    try:
        result = converter.convert(path, raises_on_error=True)
        document = result.document
        projection = project_docling_document(document)
        exported = document.export_to_markdown()
        markers = marker_facts(exported, MARKERS[label])
        required_markers = marker_facts(
            exported, REQUIRED_CONTENT_MARKERS[label]
        )
        asset_hashes = picture_hashes(document)
        return {
            "success": True,
            "elapsed_seconds": round(time.perf_counter() - started, 3),
            "status": str(result.status),
            "item_count": len(projection),
            "labels": dict(Counter(item["label"] for item in projection)),
            "table_count": len(document.tables),
            "picture_count": len(document.pictures),
            "extractable_picture_count": len(asset_hashes),
            "provenance_item_count": sum(bool(item["pages"]) for item in projection),
            "markers_complete": markers["complete"],
            "marker_order_preserved": markers["order_preserved"],
            "required_content_complete": required_markers["complete"],
            "projection_sha256": stable_json_hash(projection),
            "asset_sha256": asset_hashes,
        }
    except Exception as error:
        return {
            "success": False,
            "elapsed_seconds": round(time.perf_counter() - started, 3),
            "error_type": type(error).__name__,
            "markers_complete": False,
            "marker_order_preserved": False,
            "required_content_complete": False,
            "projection_sha256": None,
            "asset_sha256": [],
            "table_count": 0,
            "picture_count": 0,
        }


def project_docling_document(document: Any) -> list[dict[str, Any]]:
    projection: list[dict[str, Any]] = []
    for item, level in document.iterate_items():
        text = canonical_text(getattr(item, "text", ""))
        pages: list[int] = []
        boxes: list[list[float]] = []
        for provenance in getattr(item, "prov", ()) or ():
            page_number = getattr(provenance, "page_no", None)
            if isinstance(page_number, int):
                pages.append(page_number)
            bbox = getattr(provenance, "bbox", None)
            if bbox is not None:
                coordinates = [
                    getattr(bbox, name, None) for name in ("l", "t", "r", "b")
                ]
                if all(isinstance(value, (int, float)) for value in coordinates):
                    boxes.append([round(float(value), 3) for value in coordinates])
        label = getattr(getattr(item, "label", None), "value", None)
        projection.append(
            {
                "ordinal": len(projection),
                "level": int(level),
                "label": str(label or type(item).__name__),
                "text_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(),
                "text_length": len(text),
                "pages": sorted(set(pages)),
                "boxes": boxes,
            }
        )
    return projection


def picture_hashes(document: Any) -> list[str]:
    hashes: list[str] = []
    for picture in document.pictures:
        try:
            image = picture.get_image(document)
            if image is None:
                continue
            output = BytesIO()
            image.save(output, format="PNG")
            hashes.append(hashlib.sha256(output.getvalue()).hexdigest())
        except Exception:
            continue
    return sorted(hashes)


def evaluate_langchain_adapter(corpus: dict[str, Path]) -> dict[str, Any]:
    from docling.chunking import HierarchicalChunker, HybridChunker
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import PdfPipelineOptions
    from docling.document_converter import DocumentConverter, PdfFormatOption
    from langchain_docling import DoclingLoader
    from langchain_docling.loader import ExportType

    pdf_options = PdfPipelineOptions()
    pdf_options.do_ocr = True
    converter = DocumentConverter(
        format_options={
            InputFormat.PDF: PdfFormatOption(pipeline_options=pdf_options)
        }
    )
    hierarchical_chunker = HierarchicalChunker()
    hybrid_chunker = HybridChunker()
    formats: dict[str, Any] = {}
    for label, path in corpus.items():
        modes: dict[str, Any] = {}
        for name, export_type, chunker in (
            ("markdown", ExportType.MARKDOWN, hierarchical_chunker),
            (
                "doc_chunks_hierarchical",
                ExportType.DOC_CHUNKS,
                hierarchical_chunker,
            ),
            ("doc_chunks_hybrid", ExportType.DOC_CHUNKS, hybrid_chunker),
        ):
            first = load_langchain_documents(
                DoclingLoader,
                converter,
                path,
                export_type,
                chunker,
                MARKERS[label],
                REQUIRED_CONTENT_MARKERS[label],
            )
            second = load_langchain_documents(
                DoclingLoader,
                converter,
                path,
                export_type,
                chunker,
                MARKERS[label],
                REQUIRED_CONTENT_MARKERS[label],
            )
            modes[name] = {
                **first,
                "warm_elapsed_seconds": second.get("elapsed_seconds"),
                "deterministic": (
                    first["success"]
                    and second["success"]
                    and first["projection_sha256"]
                    == second["projection_sha256"]
                ),
            }
        formats[label] = modes
    return {"formats": formats}


def load_langchain_documents(
    loader_class: Any,
    converter: Any,
    path: Path,
    export_type: Any,
    chunker: Any,
    markers: tuple[str, ...],
    required_markers: tuple[str, ...],
) -> dict[str, Any]:
    started = time.perf_counter()
    try:
        loader = loader_class(
            file_path=str(path),
            converter=converter,
            export_type=export_type,
            chunker=chunker,
        )
        documents = loader.load()
        combined = "\n".join(item.page_content for item in documents)
        projection = [
            {
                "ordinal": ordinal,
                "content_sha256": hashlib.sha256(
                    canonical_text(document.page_content).encode("utf-8")
                ).hexdigest(),
                "content_length": len(canonical_text(document.page_content)),
                "metadata_keys": sorted(str(key) for key in document.metadata),
                "metadata_sha256": stable_json_hash(safe_json(document.metadata)),
            }
            for ordinal, document in enumerate(documents)
        ]
        facts = marker_facts(combined, markers)
        required_facts = marker_facts(combined, required_markers)
        return {
            "success": True,
            "document_count": len(documents),
            "elapsed_seconds": round(time.perf_counter() - started, 3),
            "markers_complete": facts["complete"],
            "marker_order_preserved": facts["order_preserved"],
            "required_content_complete": required_facts["complete"],
            "metadata_keys": sorted(
                {
                    str(key)
                    for document in documents
                    for key in document.metadata
                }
            ),
            "projection_sha256": stable_json_hash(projection),
        }
    except Exception as error:
        return {
            "success": False,
            "document_count": 0,
            "elapsed_seconds": round(time.perf_counter() - started, 3),
            "markers_complete": False,
            "marker_order_preserved": False,
            "required_content_complete": False,
            "metadata_keys": [],
            "projection_sha256": None,
            "error_type": type(error).__name__,
        }


def run_current_worker(
    project_python: Path, corpus: dict[str, Path]
) -> dict[str, Any]:
    completed = subprocess.run(
        (
            str(project_python),
            str(Path(__file__).resolve()),
            "--current-worker",
            str(next(iter(corpus.values())).parent),
        ),
        check=False,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        env={**os.environ, "PYTHONPATH": "src:."},
    )
    if completed.returncode != 0:
        return {
            "success": False,
            "returncode": completed.returncode,
            "error": "current_parser_worker_failed",
        }
    try:
        return json.loads(completed.stdout)
    except json.JSONDecodeError:
        return {
            "success": False,
            "returncode": completed.returncode,
            "error": "current_parser_worker_invalid_json",
        }


def evaluate_current_parser(corpus_root: Path) -> dict[str, Any]:
    from rag_kb.adapters.parser.langchain_unstructured import (
        partition_with_unstructured,
    )
    from rag_kb.domain import (
        AdmissionLimits,
        FileAdmissionError,
        ParserExecutionError,
        ParserLimits,
        ParserSource,
    )
    from rag_kb.services.admission import FileAdmissionService

    formats: dict[str, Any] = {}
    for label in MARKERS:
        path = corpus_root / CORPUS_FILENAMES[label]
        if not path.is_file():
            formats[label] = {"admitted": False, "error": "fixture_missing"}
            continue
        media_type = CURRENT_MEDIA_TYPES.get(
            path.suffix.lower(), "application/octet-stream"
        )
        try:
            with path.open("rb") as source:
                admitted = FileAdmissionService(AdmissionLimits()).validate(
                    source,
                    original_filename=path.name,
                    media_type=media_type,
                )
        except FileAdmissionError as error:
            formats[label] = {
                "admitted": False,
                "error_code": error.code.value,
            }
            continue
        try:
            parsed = partition_with_unstructured(
                ParserSource(path.name, admitted.media_type, path.read_bytes()),
                ParserLimits(),
            )
            combined = "\n".join(item.text for item in parsed.elements)
            facts = marker_facts(combined, MARKERS[label])
            formats[label] = {
                "admitted": True,
                "parsed": True,
                "element_count": len(parsed.elements),
                "markers_complete": facts["complete"],
                "projection_sha256": stable_json_hash(
                    [
                        {
                            "ordinal": item.ordinal,
                            "category": item.category,
                            "text_sha256": hashlib.sha256(
                                canonical_text(item.text).encode("utf-8")
                            ).hexdigest(),
                            "source_location": item.source_location,
                        }
                        for item in parsed.elements
                    ]
                ),
            }
        except ParserExecutionError as error:
            formats[label] = {
                "admitted": True,
                "parsed": False,
                "error_code": error.code.value,
            }
    return {
        "success": True,
        "versions": {
            "unstructured": package_version("unstructured"),
            "langchain_unstructured": package_version("langchain-unstructured"),
        },
        "formats": formats,
    }


def marker_facts(content: str, markers: tuple[str, ...]) -> dict[str, bool]:
    positions = [content.find(marker) for marker in markers]
    return {
        "complete": all(position >= 0 for position in positions),
        "order_preserved": all(
            left < right
            for left, right in zip(positions, positions[1:])
            if left >= 0 and right >= 0
        ),
    }


def safe_json(value: Any) -> Any:
    try:
        json.dumps(value, allow_nan=False)
        return value
    except (TypeError, ValueError):
        if isinstance(value, dict):
            return {
                str(key): safe_json(item)
                for key, item in sorted(value.items(), key=lambda pair: str(pair[0]))
            }
        if isinstance(value, (list, tuple)):
            return [safe_json(item) for item in value]
        if isinstance(value, Path):
            return value.name
        return {
            "type": f"{type(value).__module__}.{type(value).__qualname__}"
        }


def stable_json_hash(value: Any) -> str:
    return hashlib.sha256(
        json.dumps(
            value,
            ensure_ascii=False,
            allow_nan=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
    ).hexdigest()


def canonical_text(value: Any) -> str:
    return str(value or "").replace("\r\n", "\n").replace("\r", "\n").strip()


def package_version(name: str) -> str | None:
    try:
        return version(name)
    except PackageNotFoundError:
        return None


def directory_size(root: Path) -> int:
    return sum(
        path.stat().st_size
        for path in root.rglob("*")
        if path.is_file() and not path.is_symlink()
    )


def peak_rss_bytes() -> int:
    observed = int(resource.getrusage(resource.RUSAGE_SELF).ru_maxrss)
    return observed if sys.platform == "darwin" else observed * 1024


if __name__ == "__main__":
    raise SystemExit(main())
